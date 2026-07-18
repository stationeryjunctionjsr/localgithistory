"""
Pincode Master Data Generator

Reads Pincodes.xlsx from the backend root directory and generates
the pincode_master.json file used by the /api/pincodes endpoints.

Usage:
    cd "C:\\Ecommerce app\backend"
    python app/scripts/generate_pincode_data.py
"""

import json
import logging
import os

import openpyxl

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


def generate():
    # Resolve paths relative to the backend root
    script_dir = os.path.dirname(os.path.abspath(__file__))
    backend_dir = os.path.abspath(os.path.join(script_dir, "..", ".."))
    xlsx_path = os.path.join(backend_dir, "Pincodes.xlsx")
    output_path = os.path.join(backend_dir, "app", "data", "pincode_master.json")

    if not os.path.exists(xlsx_path):
        logger.error("ERROR: %s not found!", xlsx_path)
        return

    logger.info("Reading %s...", xlsx_path)
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    data = {}
    skipped = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        pincode = str(row[0].value).strip() if row[0].value else ""
        district = str(row[1].value).strip() if row[1].value else ""
        state = str(row[2].value).strip() if row[2].value else ""

        if not pincode or not district or not state:
            skipped += 1
            continue

        if state not in data:
            data[state] = {}
        if district not in data[state]:
            data[state][district] = []
        if pincode not in data[state][district]:
            data[state][district].append(pincode)

    # Sort everything for consistent output
    sorted_data = {}
    for state in sorted(data.keys()):
        sorted_data[state] = {}
        for district in sorted(data[state].keys()):
            sorted_data[state][district] = sorted(data[state][district])

    # Write output
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w") as f:
        json.dump(sorted_data, f)

    total_states = len(sorted_data)
    total_districts = sum(len(v) for v in sorted_data.values())
    total_pincodes = sum(len(p) for d in sorted_data.values() for p in d.values())

    logger.info("Done! Written to %s", output_path)
    logger.info("  States:    %s", total_states)
    logger.info("  Districts: %s", total_districts)
    logger.info("  Pincodes:  %s", total_pincodes)
    if skipped:
        logger.warning("  Skipped:   %s rows (missing data)", skipped)


if __name__ == "__main__":
    generate()
